from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from eval.inject_errors import (
    CHANGED_PROJECT_TITLE,
    IDENTITY_TEXT,
    PROJECT_TITLE,
    inject_errors,
)


def _all_text(doc) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.parametrize("error_number", range(1, 11))
def test_each_injector_changes_expected_structure(tmp_path: Path, error_number: int):
    source = PROJECT_ROOT / "eval" / "samples" / "sample_proposal.docx"
    output = tmp_path / f"error_{error_number}.docx"
    inject_errors(source, [error_number], output)
    doc = Document(output)
    text = _all_text(doc)

    if error_number == 1:
        assert "三、研究方法与技术路线" not in text
    elif error_number == 2:
        page_breaks = doc.element.body.xpath(".//w:br[@w:type='page']")
        assert len(page_breaks) >= 8
    elif error_number == 3:
        assert any("合计" in row.cells[1].text and "25.0" in _all_text_cells(row) for table in doc.tables for row in table.rows)
    elif error_number == 4:
        assert any("管理费" in _all_text_cells(row) and "6.0" in _all_text_cells(row) for table in doc.tables for row in table.rows)
    elif error_number == 5:
        assert IDENTITY_TEXT in text
    elif error_number == 6:
        assert not any(p.text.strip().startswith("图2 ") for p in doc.paragraphs)
        assert "如图2所示" in text
    elif error_number == 7:
        assert text.index("三、研究方法与技术路线") < text.index("二、研究目标与内容")
    elif error_number == 8:
        assert CHANGED_PROJECT_TITLE in text
        assert PROJECT_TITLE in text
    elif error_number == 9:
        assert "研究期限：2026年9月 — 2025年8月" in text
    elif error_number == 10:
        assert "签 字 盖 章 页" not in text
        assert "单位盖章" not in text

    assert output.with_suffix(".ground_truth.json").exists()


def _all_text_cells(row) -> str:
    return " ".join(cell.text for cell in row.cells)
