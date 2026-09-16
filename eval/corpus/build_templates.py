"""dachuang / nsfc 规则集“干净样本模板”生成器（python eval/corpus/build_templates.py [--out DIR] [--only dachuang|nsfc]），产物默认写入 eval/corpus/templates/。
产物：dachuang_clean.docx（申报书，≤8 页 / ≤6000 字 / 12 个月）、nsfc_clean.docx（申请书，由 campus 模板派生：48 个月、申请代码、推荐信、承诺字段）。
两者必须对各自规则集“干净”（0 确定性 error/warning，由 selfcheck.py --ruleset 逐份校准）；本模块只生成，不写数据库、不调用模型。
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EVAL_DIR = PROJECT_ROOT / "eval"
sys.path.insert(0, str(EVAL_DIR))
sys.path.insert(0, str(Path(__file__).resolve().parent))

# Windows 控制台/重定向默认 cp1252：中文输出先重配 UTF-8
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

# 复用已校准的排版助手（标题字体/加粗、正文宋体 12、表格样式）
import generate_sample_docx as campus_builder  # noqa: E402

CAMPUS_TEMPLATE = EVAL_DIR / "samples" / "sample_proposal.docx"
TEMPLATE_DIR = EVAL_DIR / "corpus" / "templates"


def _set_east_asia(run, name: str) -> None:
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run._element.insert(0, rpr)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)


def _replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new, 1)
            return True
    text = paragraph.text.replace(old, new, 1)
    paragraph.clear()
    paragraph.add_run(text)
    return True


def _insert_paragraph_after(anchor, text: str, *, size: int = 14, font: str = "宋体",
                            center: bool = True):
    paragraph = anchor.insert_paragraph_before()
    # insert_paragraph_before 放在 anchor 之前；改用 XML 追加到 anchor 之后
    anchor._p.addnext(paragraph._p)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    _set_east_asia(run, font)
    return paragraph


def _insert_body_after(anchor, text: str):
    """在 anchor 之后插入正文段落（宋体 12、首行缩进、1.5 倍行距，与模板正文一致）。"""
    paragraph = _insert_paragraph_after(anchor, text, size=12, center=False)
    pf = paragraph.paragraph_format
    pf.first_line_indent = Cm(0.74)
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    return paragraph


# ── nsfc：由 campus 模板派生 ──────────────────────────────
def build_nsfc_template(campus_docx: Path, out_path: Path) -> None:
    """campus 模板 + NSFC 差异：48 个月期限、申请代码、推荐信、承诺。"""
    doc = Document(str(campus_docx))
    paragraphs = list(doc.paragraphs)

    # 1) 封面：研究期限 24 → 48 个月；补“申请代码”
    for paragraph in paragraphs:
        if "研究期限：2026年9月 — 2028年8月" in paragraph.text:
            assert _replace_in_paragraph(
                paragraph, "2028年8月", "2030年8月"
            ), "封面研究期限替换失败"
        if paragraph.text.startswith("申报单位："):
            _insert_paragraph_after(paragraph, "申请代码：F0301")

    # 2) 正文进度安排中的期限描述必须同步（date_check 会聚合多处日期）
    for paragraph in paragraphs:
        if "本课题研究期限为2026年9月至2028年8月" in paragraph.text:
            _replace_in_paragraph(paragraph, "2028年8月", "2030年8月")
            _replace_in_paragraph(paragraph, "共计24个月", "共计48个月")

    # 3) 附件清单：补“推荐信”
    attachment_anchor = None
    for paragraph in paragraphs:
        if paragraph.text.startswith("附件：依托单位推荐意见"):
            attachment_anchor = paragraph
            break
    assert attachment_anchor is not None, "未找到附件清单锚点"
    _insert_paragraph_after(attachment_anchor, "附件：专家推荐信", size=12, center=False)

    # 4) 签字盖章页：补“承诺”关键字（nsfc signature_page keywords 含 承诺）
    signature_anchor = None
    for paragraph in paragraphs:
        if "依托单位审核意见" in paragraph.text:
            signature_anchor = paragraph
            break
    assert signature_anchor is not None, "未找到签字页锚点"
    _insert_paragraph_after(
        signature_anchor, "本人承诺：以上信息真实、准确，愿承担相应责任。",
        size=14, center=False,
    )

    # 5) NSFC 的分节字数下限高于校级模板：只在本派生模板内补充内容
    #    （立项依据 ≥800 字；`研究内容|研究目标` 实际匹配到 2.2 研究内容，需 ≥500 字）
    basis_anchor = None
    for paragraph in paragraphs:
        if paragraph.text.startswith("从实际应用价值来看"):
            basis_anchor = paragraph
            break
    assert basis_anchor is not None, "未找到立项依据末段锚点"
    _insert_body_after(basis_anchor,
        "与既有研究相比，本项目的切入点在于把动态图结构与多尺度时间建模统一到同一框架中："
        "一方面通过门控机制让邻接矩阵随突发事件自适应更新，另一方面用分层时间卷积刻画"
        "分钟级波动与日级趋势。这种设计既保留了图方法对路网拓扑的建模能力，"
        "又避免了纯注意力结构在大规模路网上的计算开销，具备较好的理论价值与推广空间。"
        "在数据层面，项目将同时使用公开基准数据集与本地城市交通数据，"
        "通过跨城市迁移实验检验模型的泛化能力；在评价层面，除常规误差指标外，"
        "还将引入高峰时段误差与长时预测误差两类指标，以更贴近交通管理部门的实际需求。"
        "预期研究成果可为城市交通信号优化、出行诱导与应急疏导提供可复用的预测底座，"
        "并为后续将方法推广到物流配送、应急疏散等相似时空预测任务提供方法参考。"
    )

    content_anchor = None
    for paragraph in paragraphs:
        if paragraph.text.startswith("研究内容四：大规模实验验证"):
            content_anchor = paragraph
            break
    assert content_anchor is not None, "未找到研究内容末段锚点"
    _insert_body_after(content_anchor,
        "上述四项研究内容相互支撑：动态拓扑建模提供结构基础，多尺度融合提升预测精度，"
        "轻量化推理保证工程可用性，而大规模实验验证则检验整体方案在真实数据上的有效性，"
        "共同构成从方法创新到应用验证的完整研究链条。"
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# ── dachuang：紧凑版申报书（≤8 页 / ≤6000 字 / 12 个月）──
def build_dachuang_template(out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # ── 封面 ──
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("大学生创新创业训练计划申报书")
    run.font.name = "黑体"
    run.font.size = Pt(24)
    run.font.bold = True
    _set_east_asia(run, "黑体")
    doc.add_paragraph()

    info_lines = [
        "项目名称：基于深度学习的校园快递路径优化研究",
        "项目负责人：李四",
        "指导教师：王五",
        "所属学院：计算机科学与技术学院",
        "填表日期：2026年6月30日",
        "联系电话：13900139000",
        "电子邮箱：lisi@university.edu.cn",
        "预算总额：50000元",
        # 研究期限放最后：date_check 会聚合其前后各 2 个文本块中的日期，
        # 相邻块不得再出现别的日期，否则期限会被算成 15 个月。
        "研究期限：2026年9月 — 2027年8月",
    ]
    for line in info_lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(line)
        run.font.name = "宋体"
        run.font.size = Pt(14)
        _set_east_asia(run, "宋体")

    doc.add_page_break()

    # ── 一、立项依据（≥200 字）──
    campus_builder.add_heading_styled(doc, "一、立项依据", level=1)
    campus_builder.add_body_para(doc,
        "校园快递最后一公里配送长期依赖人工经验，高峰时段取件排队时间长、配送路线重复，"
        "造成人力浪费与体验下降。随着校园快递量逐年增长，仅靠增加人手难以持续改善。"
    )
    campus_builder.add_body_para(doc,
        "现有路径优化研究多面向城市配送场景，缺少对校园楼宇分布、课程作息与人流潮汐的建模，"
        "直接套用效果有限。本项目拟结合校园时空数据与启发式算法，构建可落地的路径优化方案，"
        "为解决校园配送效率问题提供可复制的技术路径。"
    )
    campus_builder.add_body_para(doc,
        "从育人角度看，本项目覆盖数据采集、算法设计、系统实现与效果评估全过程，"
        "能够让团队成员在真实问题中锻炼工程能力与协作能力，具备良好的训练价值。"
    )

    # ── 二、研究内容（≥200 字）──
    campus_builder.add_heading_styled(doc, "二、研究内容", level=1)
    campus_builder.add_body_para(doc,
        "研究内容一：校园配送数据采集与建模。整合快递驿站历史取件记录、楼宇分布与课程作息，"
        "构建带时间窗的配送需求模型，刻画不同时段的取件潮汐特征。"
    )
    campus_builder.add_body_para(doc,
        "研究内容二：路径优化算法设计与实现。在经典车辆路径问题基础上引入时间窗与载量约束，"
        "设计改进的启发式求解算法，并与贪心、遗传算法等基线方法进行对比实验。"
    )
    campus_builder.add_body_para(doc,
        "研究内容三：原型系统与效果评估。开发轻量级调度原型系统，在校内小范围试运行，"
        "以配送里程、平均等待时长和人力投入为指标评估优化效果。"
    )

    # ── 三、创新点（≥100 字）──
    campus_builder.add_heading_styled(doc, "三、创新点", level=1)
    campus_builder.add_body_para(doc,
        "创新点一：将课程作息与楼宇分布引入时间窗建模，使路径优化贴合校园真实人流规律，"
        "相比通用城市配送模型更贴近校园场景，可直接用于校内小范围调度验证。"
    )
    campus_builder.add_body_para(doc,
        "创新点二：提出面向小规模多约束场景的轻量启发式算法，降低对算力与数据的依赖，"
        "便于在学生团队条件下实现与验证，并给出与经典算法的量化对比结论。"
    )

    # ── 四、进度安排 ──
    campus_builder.add_heading_styled(doc, "四、进度安排", level=1)
    campus_builder.add_body_para(doc,
        "项目周期 12 个月，分四个阶段推进，各阶段任务与时间安排如下表所示。"
    )
    campus_builder.create_table(doc, ["阶段", "时间", "主要任务"], [
        ["第一阶段", "2026年9月 - 2026年11月", "需求调研、数据采集与清洗"],
        ["第二阶段", "2026年12月 - 2027年2月", "需求建模与算法设计"],
        ["第三阶段", "2027年3月 - 2027年5月", "算法实现与对比实验"],
        ["第四阶段", "2027年6月 - 2027年8月", "原型系统试运行与结题报告"],
    ])

    # ── 五、经费预算（合计=分项之和；单位用“元”以便与封面预算总额同口径）──
    campus_builder.add_heading_styled(doc, "五、经费预算", level=1)
    campus_builder.add_body_para(doc,
        "本项目申请经费合计 50000 元，全部为小额实验与调研支出，明细如下表所示。"
    )
    campus_builder.create_table(doc, ["序号", "经费科目", "金额（元）", "计算依据及说明"], [
        ["1", "材料费", "12000", "实验耗材与数据存储介质"],
        ["2", "测试化验加工费", "10000", "云服务器租用与数据标注"],
        ["3", "差旅费", "10000", "校内调研与学术会议"],
        ["4", "劳务费", "15000", "学生助研津贴"],
        ["5", "管理费", "3000", "按 6% 计提的管理费"],
        ["", "合计", "50000", ""],
    ])
    campus_builder.add_body_para(doc, "注：管理费占比 6%，未超过规定上限。")

    # ── 六、预期成果 ──
    campus_builder.add_heading_styled(doc, "六、预期成果", level=1)
    campus_builder.add_body_para(doc,
        "预期形成校园配送路径优化算法一套、调度原型系统一个，完成结题报告一份；"
        "在试运行阶段形成可量化的效率提升数据，并整理形成 1 篇研究论文或技术报告的初稿。"
    )

    # ── 七、附件清单 ──
    campus_builder.add_heading_styled(doc, "七、附件清单", level=1)
    for text in ["附件：指导教师意见（签字）", "附件：学院推荐意见（盖章）"]:
        campus_builder.add_body_para(doc, text)

    # ── 签字盖章页 ──
    doc.add_page_break()
    for _ in range(4):
        doc.add_paragraph()
    signature_title = doc.add_paragraph()
    signature_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = signature_title.add_run("指导教师意见与学院审核")
    run.font.name = "黑体"
    run.font.size = Pt(16)
    run.font.bold = True
    _set_east_asia(run, "黑体")
    doc.add_paragraph()
    for text in [
        "指导教师签字：________________    日期：____________",
        "",
        "学院盖章：________________        日期：____________",
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        run.font.name = "宋体"
        run.font.size = Pt(14)
        _set_east_asia(run, "宋体")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def ensure_campus_template() -> Path:
    if not CAMPUS_TEMPLATE.exists():
        campus_builder.generate_sample(str(CAMPUS_TEMPLATE))
    return CAMPUS_TEMPLATE


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=str(TEMPLATE_DIR), help="模板输出目录")
    parser.add_argument("--only", choices=["dachuang", "nsfc"], help="只生成其中一个")
    args = parser.parse_args()

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    campus = ensure_campus_template()

    if args.only != "nsfc":
        target = out_dir / "dachuang_clean.docx"
        build_dachuang_template(target)
        print(f"[OK] dachuang 模板：{target}")
    if args.only != "dachuang":
        target = out_dir / "nsfc_clean.docx"
        build_nsfc_template(campus, target)
        print(f"[OK] nsfc 模板：{target}")
    print("\n下一步校准（各自应为 0 确定性 error/warning）：")
    print("  python eval/corpus/selfcheck.py --docx "
          f"{out_dir / 'dachuang_clean.docx'} --ruleset dachuang")
    print("  python eval/corpus/selfcheck.py --docx "
          f"{out_dir / 'nsfc_clean.docx'} --ruleset nsfc")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
