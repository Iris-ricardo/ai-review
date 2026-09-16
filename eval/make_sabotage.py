"""Generate a sabotaged version of sample_proposal.docx for M2/M3 testing.

Sabotage: budget table total 30.0 → 25.0 (triggers budget_check error). Usage: cd eval && python make_sabotage.py → eval/samples/sabotaged_budget.docx
"""
from pathlib import Path
from docx import Document

SAMPLES_DIR = Path(__file__).parent / "samples"
SRC = SAMPLES_DIR / "sample_proposal.docx"
DST = SAMPLES_DIR / "sabotaged_budget.docx"


def make_sabotaged(output_path: str | Path | None = None):
    if not SRC.exists():
        raise FileNotFoundError(f"Source not found: {SRC}\nRun generate_sample_docx.py first.")

    dst = Path(output_path) if output_path is not None else DST
    doc = Document(str(SRC))

    sabotaged = False
    for table in doc.tables:
        for row in table.rows:
            cells_text = [c.text.strip() for c in row.cells]
            # Find the total row (contains "合计")
            if "合计" in str(cells_text):
                for cell in row.cells:
                    if "30.0" in cell.text:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                if "30.0" in run.text:
                                    run.text = run.text.replace("30.0", "25.0")
                                    sabotaged = True
                                    print(f"  Budget total: 30.0 → 25.0")
    if not sabotaged:
        raise RuntimeError("Could not find budget total cell (30.0) in sample_proposal.docx")

    doc.save(str(dst))
    print(f"Sabotaged DOCX saved to: {dst}")
    return dst


if __name__ == "__main__":
    make_sabotaged()
