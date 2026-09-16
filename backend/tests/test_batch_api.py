from __future__ import annotations

from io import BytesIO
from pathlib import Path
import time

from starlette.testclient import TestClient
from openpyxl import load_workbook

from app.api import routes
from app.core.config import get_settings
from app.main import app


def test_batch_three_files_end_to_end(tmp_path: Path, monkeypatch):
    sample_pdf = Path(__file__).resolve().parents[2] / "eval" / "samples" / "sample_proposal.pdf"
    assert sample_pdf.exists()
    content = sample_pdf.read_bytes()

    monkeypatch.setenv("LLM_API_KEY", "")
    get_settings.cache_clear()
    monkeypatch.setattr(routes.settings, "UPLOAD_DIR", str(tmp_path / "uploads"))
    monkeypatch.setattr(routes.settings, "OUTPUT_DIR", str(tmp_path / "outputs"))

    client = TestClient(app)
    response = client.post(
        "/api/v1/batches",
        # R08：批量流程用例不涉及 AI，显式关闭以免与"授权 + 部署外发开关"耦合
        data={"ruleset_id": "campus_general_v1", "use_ai": "false"},
        files=[
            ("files", (f"sample-{index}.pdf", content, "application/pdf"))
            for index in range(1, 4)
        ],
    )
    assert response.status_code == 200, response.text
    batch_id = response.json()["batch_id"]

    deadline = time.monotonic() + 30
    while True:
        status = client.get(f"/api/v1/batches/{batch_id}")
        assert status.status_code == 200
        payload = status.json()
        if payload["done_count"] == payload["total"]:
            break
        if time.monotonic() >= deadline:
            raise AssertionError(f"Batch {batch_id} did not finish before timeout")
        time.sleep(0.05)
    assert payload["total"] == 3
    assert payload["done_count"] == 3
    assert len(payload["files"]) == 3
    assert all(item["status"] == "done" for item in payload["files"])
    assert all(item["review_id"] for item in payload["files"])

    summary = client.get(f"/api/v1/batches/{batch_id}/summary.xlsx")
    assert summary.status_code == 200
    workbook = load_workbook(BytesIO(summary.content), read_only=True)
    try:
        sheet = workbook["审查汇总"]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert rows[0] == ("文件名", "结论", "error数", "warning数", "info数", "审查时间")
    assert len(rows) == 4
